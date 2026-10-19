"""Regressões do leitor local: tipos, estrutura, privacidade e falhas explícitas."""
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from src.tools.file_analysis import capabilities, ingest, inspect
from src.tools.file_analysis.contracts import FileRejected, MAX_BYTES, detect, inspect_archive
from src.tools.file_analysis.extractors import extract

try:
    from PIL import Image
except ImportError:
    Image = None
try:
    from docx import Document
except ImportError:
    Document = None
try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from reportlab.pdfgen.canvas import Canvas
except ImportError:
    Canvas = None


class ContractsTests(unittest.TestCase):
    def test_capabilities_do_not_pretend_vision_is_installed(self):
        result = capabilities()
        self.assertEqual(result["name"], "files.inspect")
        self.assertEqual(result["visual_understanding"], "not_integrated")
        self.assertEqual(result["ocr"], "not_integrated")

    def test_bad_signature_path_large_file_and_wrong_encoding(self):
        bad = [
            (b"not a pdf", "report.pdf"), (b"garbage", "photo.png"),
            (b"abc", "../../private.txt"), (b"\xff\xfe", "a.txt"),
            (b"PKgarbage", "a.docx"), (b"x", "script.exe"),
            (b"\x00secret", "note.md"), (b"a" * (MAX_BYTES + 1), "huge.txt"),
        ]
        for data, name in bad:
            with self.subTest(name=name), self.assertRaises(FileRejected):
                detect(data, name)

    def test_zip_bomb_rejected_before_office_parser(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as target:
            target.writestr("[Content_Types].xml", b"<Types/>")
            target.writestr("word/document.xml", b"A" * (1024 * 1024))
        with self.assertRaises(FileRejected):
            inspect_archive(buf.getvalue())

    def test_csv_preserves_columns(self):
        data = b"produto;quantidade\ncafe;3\narroz;5\n"
        result = extract(data, "compras.csv")
        self.assertEqual(result["structure"][0]["rows"][1], ["cafe", "3"])


class StorageTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name) / "data"

    def test_private_ingestion_and_bounded_worker(self):
        receipt = ingest("Projeto: menor IA".encode(), "planejamento.md", root=self.root)
        self.assertEqual(receipt["status"], "executed")
        file_id = receipt["file"]["file_id"]
        self.assertRegex(file_id, r"^[a-f0-9]{32}$")
        private = self.root / "media" / "file_analysis"
        self.assertEqual((private / (file_id + ".blob")).stat().st_mode & 0o777, 0o600)
        result = inspect(file_id, root=self.root)
        self.assertEqual(result["status"], "executed", result)
        self.assertEqual(result["content"][0]["text"], "Projeto: menor IA")
        self.assertTrue(result["validation"]["format_verified"])
        self.assertFalse(result["validation"]["semantic_verified"])
        self.assertEqual(result["evidence"][0]["kind"], "text_decode")

    def test_private_integrity_and_no_arbitrary_paths(self):
        receipt = ingest(b"arquivo privado", "memo.txt", root=self.root)
        token = receipt["file"]["file_id"]
        with self.assertRaises(FileRejected):
            inspect("../../etc/passwd", root=self.root)
        private = self.root / "media" / "file_analysis"
        (private / (token + ".blob")).write_bytes(b"alterado")
        with self.assertRaises(FileRejected):
            inspect(token, root=self.root)

    def test_private_symlink_rejected(self):
        if not hasattr(os, "symlink"):
            self.skipTest("Symlink indisponível")
        receipt = ingest(b"texto", "memo.txt", root=self.root)
        token = receipt["file"]["file_id"]
        private = self.root / "media" / "file_analysis"
        blob = private / (token + ".blob")
        blob.unlink()
        blob.symlink_to(private / (token + ".json"))
        with self.assertRaises(FileRejected):
            inspect(token, root=self.root)

    def test_worker_environment_excludes_secrets_and_timeout_fails_closed(self):
        receipt = ingest(b"abc", "memo.txt", root=self.root)
        with patch.dict(os.environ, {"MENOR_IMAGE_API_KEY": "private-token"}):
            with patch("src.tools.file_analysis.service.subprocess.run") as run:
                run.side_effect = __import__("subprocess").TimeoutExpired("worker", 1)
                result = inspect(receipt["file"]["file_id"], root=self.root)
                self.assertEqual(result["status"], "failed")
                self.assertEqual(result["error"]["code"], "timeout")
                self.assertNotIn("MENOR_IMAGE_API_KEY", run.call_args.kwargs["env"])


@unittest.skipUnless(Image, "Instale Pillow para testes de imagem")
class ImageTests(unittest.TestCase):
    def test_image_dimensions_and_no_false_visual_claim(self):
        buf = io.BytesIO()
        Image.new("RGB", (32, 32), "purple").save(buf, "PNG")
        result = extract(buf.getvalue(), "sprite.png")
        self.assertEqual(result["structure"][0]["width"], 32)
        self.assertEqual(result["structure"][0]["height"], 32)
        self.assertIn("não comprovam conteúdo", result["limitations"][0])

    def test_invalid_image_fails(self):
        with self.assertRaises(FileRejected):
            extract(b"\x89PNG\r\n\x1a\nINVALID", "fake.png")


@unittest.skipUnless(Document, "Instale python-docx para testar DOCX")
class DocxTests(unittest.TestCase):
    def test_headings_and_tables(self):
        doc = Document()
        doc.add_heading("Introdução", 1)
        doc.add_paragraph("Texto extraído corretamente.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Preço"
        table.cell(1, 0).text = "Livro"
        table.cell(1, 1).text = "25"
        out = io.BytesIO()
        doc.save(out)
        result = extract(out.getvalue(), "teste.docx")
        self.assertEqual(result["structure"][0]["heading"], True)
        table_result = next(item for item in result["structure"] if item["type"] == "table")
        self.assertEqual(table_result["rows"][1], ["Livro", "25"])


@unittest.skipUnless(Workbook, "Instale openpyxl para testar XLSX")
class XlsxTests(unittest.TestCase):
    def test_sheet_names_cells_and_cached_formula_limitation(self):
        wb = Workbook()
        wb.active.title = "Vendas"
        wb.active.append(["Mês", "Total"])
        wb.active.append(["Janeiro", 150])
        wb.create_sheet("Resumo")
        out = io.BytesIO()
        wb.save(out)
        result = extract(out.getvalue(), "planilha.xlsx")
        self.assertEqual(result["metadata"]["sheet_names"], ["Vendas", "Resumo"])
        self.assertEqual(result["structure"][0]["rows"][1]["cells"][1]["value"], 150)
        self.assertIn("cache", result["limitations"][0])

    def test_xlsx_end_to_end_private_worker(self):
        wb = Workbook()
        wb.active.title = "Inventário"
        wb.active["A1"] = "Código"
        wb.active["B1"] = "Quantidade"
        wb.active["A2"] = "ABC-1"
        wb.active["B2"] = 7
        out = io.BytesIO()
        wb.save(out)
        with tempfile.TemporaryDirectory() as tmp:
            stored = ingest(out.getvalue(), "inventario.xlsx", root=Path(tmp) / "data")
            receipt = inspect(stored["file"]["file_id"], root=Path(tmp) / "data")
        self.assertEqual(receipt["status"], "executed", receipt)
        self.assertEqual(receipt["content"][0]["name"], "Inventário")
        self.assertEqual(receipt["content"][0]["rows"][1]["cells"][1]["value"], 7)

    def test_mismatched_office_extension(self):
        wb = Workbook()
        out = io.BytesIO()
        wb.save(out)
        with self.assertRaises(FileRejected):
            detect(out.getvalue(), "arquivo.docx")


@unittest.skipUnless(PdfReader and Canvas, "Instale pypdf e reportlab para testes PDF")
class PdfTests(unittest.TestCase):
    def test_pdf_page_and_text_origin(self):
        out = io.BytesIO()
        pdf = Canvas(out)
        pdf.setTitle("Relatório")
        pdf.drawString(72, 720, "Relatorio do Menor")
        pdf.showPage()
        pdf.drawString(72, 720, "Pagina dois")
        pdf.save()
        result = extract(out.getvalue(), "teste.pdf")
        self.assertEqual(result["metadata"]["pages_total"], 2)
        self.assertIn("Relatorio do Menor", result["structure"][0]["text"])
        self.assertEqual(result["evidence"][1]["number"], 2)

    def test_corrupt_pdf_never_succeeds(self):
        with self.assertRaises(Exception):
            extract(b"%PDF-1.4\ncorrupted", "arquivo.pdf")


if __name__ == "__main__":
    unittest.main()
