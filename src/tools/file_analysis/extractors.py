"""Extratores determinísticos e limitados; não executam conteúdo dos anexos."""
import csv
import io
import warnings

from src.tools.file_analysis.contracts import (
    FileRejected, MAX_COLS, MAX_PAGES, MAX_PIXELS, MAX_ROWS, MAX_SHEETS, MAX_TEXT,
    detect,
)


def _clip(value, limit=2500):
    return str(value)[:limit] if value is not None else ""


def _text(data, fmt):
    text = data.decode("utf-8-sig")
    if fmt == "csv":
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect, strict=True)
        rows, truncated = [], False
        for index, row in enumerate(reader):
            if index >= MAX_ROWS:
                truncated = True
                break
            if len(row) > MAX_COLS:
                truncated = True
            rows.append([_clip(cell, 1000) for cell in row[:MAX_COLS]])
        return {
            "structure": [{"type": "table", "rows": rows, "row_count": len(rows)}],
            "metadata": {"rows_read": len(rows)}, "truncated": truncated,
            "evidence": [{"kind": "csv_rows", "count": len(rows)}],
            "limitations": ["Linhas e colunas longas podem ser truncadas."],
        }
    truncated = len(text) > MAX_TEXT
    return {
        "structure": [{"type": "text", "text": text[:MAX_TEXT]}],
        "metadata": {"characters": len(text)}, "truncated": truncated,
        "evidence": [{"kind": "text_decode", "encoding": "utf-8"}],
        "limitations": ["Conteúdo tratado como dados; instruções internas não são executadas."],
    }


def _pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        raise FileRejected("PDF criptografado não é suportado.")
    pages, used, truncated = [], 0, False
    for index, page in enumerate(reader.pages):
        if index >= MAX_PAGES or used >= MAX_TEXT:
            truncated = True
            break
        # pypdf não executa scripts; páginas podem conter estruturas complexas.
        extracted = page.extract_text() or ""
        remaining = MAX_TEXT - used
        limited = extracted[:min(6000, remaining)]
        used += len(limited)
        if len(extracted) > len(limited):
            truncated = True
        pages.append({"type": "page", "page": index + 1, "text": limited})
    meta = reader.metadata
    metadata = {"pages_total": len(reader.pages), "pages_read": len(pages)}
    if meta:
        for field in ("title", "author", "subject"):
            raw = getattr(meta, field, None)
            if isinstance(raw, str):
                metadata[field] = _clip(raw, 300)
    return {
        "structure": pages, "metadata": metadata, "truncated": truncated,
        "evidence": [{"kind": "pdf_page", "number": p["page"]} for p in pages],
        "limitations": ["Não há OCR nesta versão; páginas digitalizadas podem retornar texto vazio.",
                        "Layout, figuras e tabelas do PDF podem não ser preservados."],
    }


def _docx(data):
    from docx import Document

    doc = Document(io.BytesIO(data))
    items, used, truncated = [], 0, False
    for item in doc.iter_inner_content():
        if used >= MAX_TEXT or len(items) >= 700:
            truncated = True
            break
        if hasattr(item, "rows"):
            rows = []
            for row_number, row in enumerate(item.rows):
                if row_number >= MAX_ROWS:
                    truncated = True
                    break
                if len(row.cells) > MAX_COLS:
                    truncated = True
                rows.append([_clip(cell.text, 800) for cell in row.cells[:MAX_COLS]])
            items.append({"type": "table", "rows": rows})
            used += sum(len(cell) for row in rows for cell in row)
        else:
            text = item.text or ""
            available = min(2500, MAX_TEXT - used)
            if len(text) > available:
                truncated = True
            text = text[:available]
            used += len(text)
            if text.strip():
                style = _clip(item.style.name if item.style else "", 100)
                items.append({"type": "paragraph", "style": style, "text": text,
                              "heading": style.lower().startswith(("heading", "título"))})
    return {
        "structure": items, "metadata": {"blocks_read": len(items)},
        "truncated": truncated,
        "evidence": [{"kind": "docx_block", "number": i + 1,
                      "type": block["type"]} for i, block in enumerate(items)],
        "limitations": ["Elementos flutuantes, revisões e imagens incorporadas não são interpretados.",
                        "Macros e links externos não são executados."],
    }


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return _clip(value, 1000)


def _xlsx(data):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise FileRejected("Planilha excede o limite de abas.")
        sheets, truncated = [], False
        for ws in wb.worksheets:
            rows = []
            limit_rows = min(ws.max_row or MAX_ROWS, MAX_ROWS)
            limit_cols = min(ws.max_column or MAX_COLS, MAX_COLS)
            if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS:
                truncated = True
            iterator = ws.iter_rows(min_row=1, max_row=limit_rows,
                                    max_col=limit_cols, values_only=True)
            for number, row in enumerate(iterator, 1):
                cells = [{"column": get_column_letter(index), "value": _cell_value(value)}
                         for index, value in enumerate(row, 1) if value is not None]
                if cells:
                    rows.append({"number": number, "cells": cells})
            sheets.append({"type": "sheet", "name": ws.title, "rows": rows,
                           "rows_read": limit_rows, "cols_read": limit_cols})
        return {
            "structure": sheets, "metadata": {"sheet_names": list(wb.sheetnames)},
            "truncated": truncated,
            "evidence": [{"kind": "xlsx_sheet", "name": name} for name in wb.sheetnames],
            "limitations": ["Fórmulas não são executadas; valores dependem do cache salvo no XLSX.",
                            "Gráficos, macros e conexões externas não são executados nem interpretados."],
        }
    finally:
        wb.close()


def _image(data, expected):
    from PIL import Image, UnidentifiedImageError

    try:
        with Image.open(io.BytesIO(data)) as pic:
            actual = {"JPEG": "jpeg", "PNG": "png", "WEBP": "webp",
                      "GIF": "gif", "BMP": "bmp"}.get(pic.format)
            if actual != expected:
                raise FileRejected("Formato da imagem diferente da assinatura.")
            width, height = pic.size
            if not width or not height or width * height > MAX_PIXELS:
                raise FileRejected("Imagem excede resolução máxima.")
            frames = getattr(pic, "n_frames", 1)
            color_mode = pic.mode
            pic.verify()
        sprite_metrics = None
        if width <= 256 and height <= 256 and frames == 1:
            # Paleta é uma medição objetiva; não prova que há um personagem.
            with Image.open(io.BytesIO(data)) as pixels:
                rgba = pixels.convert("RGBA")
                histogram = rgba.getcolors(maxcolors=257)
                palette_count = len(histogram) if histogram is not None else None
                if histogram is not None:
                    transparent = sum(count for count, color in histogram if color[3] == 0)
                    partial_alpha = sum(count for count, color in histogram if 0 < color[3] < 255)
                else:
                    transparent, partial_alpha = None, None
                visible_bounds = rgba.getchannel("A").getbbox()
                bounds = list(visible_bounds) if transparent and visible_bounds else None
                sprite_metrics = {
                    "unique_rgba_colors": palette_count,
                    "more_than_256_colors": histogram is None,
                    "fully_transparent_pixels": transparent,
                    "partially_transparent_pixels": partial_alpha,
                    "nontransparent_bounds": bounds,
                }
    except (OSError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise FileRejected("Imagem inválida ou danificada.") from exc
    return {
        "structure": [{"type": "image", "format": actual, "width": width,
                       "height": height, "frames": frames, "mode": color_mode,
                       "sprite_metrics": sprite_metrics}],
        "metadata": {"width": width, "height": height, "format": actual},
        "truncated": False,
        "evidence": [{"kind": "image_integrity", "verified": True,
                      "dimensions": [width, height]}],
        "limitations": ["Integridade e dimensão não comprovam conteúdo, estilo ou qualidade visual.",
                        "Sem modelo multimodal operacional não há descrição da cena ou OCR."],
    }


def extract(data, filename):
    """Analisa bytes (nunca um path do usuário); dados retornam com origem explícita."""
    fmt = detect(data, filename)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", category=UserWarning)
            if fmt == "pdf":
                parsed = _pdf(data)
            elif fmt == "docx":
                parsed = _docx(data)
            elif fmt == "xlsx":
                parsed = _xlsx(data)
            elif fmt in ("png", "jpeg", "webp", "gif", "bmp"):
                parsed = _image(data, fmt)
            else:
                parsed = _text(data, fmt)
    except (FileRejected, ImportError):
        raise
    except (ValueError, TypeError, OSError, RuntimeError, KeyError, csv.Error, EOFError) as exc:
        raise FileRejected("Não foi possível interpretar o arquivo.") from exc
    return {"format": fmt, **parsed}
